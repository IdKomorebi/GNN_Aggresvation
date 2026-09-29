"""Plain white PowerPoint progress reports. Original experiment PNGs preferred.

Install python-pptx if needed. In this workspace it is staged in /tmp.
PowerPoint crops preserve the original image bytes and can be reset by the user.
"""
from pathlib import Path
import sys,json,hashlib,re,math
DEPS=Path('/tmp/group-meeting-pptx-deps')
if DEPS.exists():sys.path.insert(0,str(DEPS))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR,MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image,ImageFont

HERE=Path(__file__).resolve().parent
OUT=HERE.parent
ROOT=OUT.parents[1]
BLUE='005A9C';BLACK='000000';GRAY='666666'
FONT='微软雅黑'
MEASURE='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
DECKS=[]

def log(n):return f'DNN_Aggresvation{n}/CHANGELOG.md'
def fig(n,name,region=None,label=''):
    return {'source':f'DNN_Aggresvation{n}/figures/{name}.png','crop':region,'label':label}
def deck(no,title,scope,work):
    d={'no':no,'title':title,'scope':scope,'slides':[]};DECKS.append(d)
    add(d,'本次工作：'+title,work,kind='intro',lead=scope,sources=[])
    return d
def add(d,title,bullets,images=None,kind='figure',sources=None,notes=None,lead='',headers=None,rows=None,widths=None,diagram=None):
    images=images or []
    if not images and kind=='figure':kind='text'
    src=list(sources or [])+[x['source'] for x in images]
    d['slides'].append({'title':title,'bullets':bullets,'images':images,'kind':kind,'sources':list(dict.fromkeys(src)),
      'notes':notes or [],'lead':lead,'headers':headers,'rows':rows,'widths':widths,'diagram':diagram})

def setfont(run,size,bold=False,color=BLACK):
    f=run.font;f.name=FONT;f.size=Pt(size);f.bold=bold;f.color.rgb=RGBColor.from_string(color)
    rp=run._r.get_or_add_rPr()
    ea=rp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
    if ea is None:ea=OxmlElement('a:ea');rp.append(ea)
    ea.set('typeface',FONT)

def tx(slide,text,x,y,w,h,size=20,bold=False,color=BLACK,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align;p.line_spacing=1.10
        p.space_after=Pt(0);p.space_before=Pt(0);run=p.add_run();run.text=line;setfont(run,size,bold,color)
    return sh

def estimate_lines(s,size,width):
    f=ImageFont.truetype(MEASURE,round(size*2));cap=width*72*2
    total=0
    for line in s.split('\n'):
        current=0;n=1
        for c in line:
            wc=f.getlength(c)
            if current+wc>cap:n+=1;current=0
            current+=wc
        total+=n
    return total

def bulletbox(slide,items,x,y,w,h,size=20,gap=12):
    rendered=['• '+item.removesuffix('。') for item in items]
    while size>17:
        needed=sum(estimate_lines(t,size,w)*size*1.12 for t in rendered)+max(0,len(items)-1)*gap
        if needed<=h*72-5:break
        size-=.5
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,item in enumerate(rendered):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.line_spacing=1.12;p.space_after=Pt(gap);p.space_before=Pt(0)
        r=p.add_run();r.text=item;setfont(r,size)
    sh.name='说明文字'
    return sh

def picture(slide,spec,x,y,w,h):
    path=ROOT/spec['source']
    with Image.open(path) as im:iw,ih=im.size
    l,t,r,b=spec['crop'] or (0,0,1,1)
    assert 0<=l<r<=1 and 0<=t<b<=1
    aspect=iw*(r-l)/(ih*(b-t));pw=min(w,h*aspect);ph=pw/aspect
    pic=slide.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
    pic.crop_left=l;pic.crop_top=t;pic.crop_right=1-r;pic.crop_bottom=1-b
    pic.name=spec['source']+(' [原始数据Matplotlib子图]' if spec.get('derived') else (' [原图子图裁剪]' if spec['crop'] else ' [原图]'))
    pic._element.nvPicPr.cNvPr.set('descr',spec['label'] or spec['source'])
    return pic

def table(slide,s,x=.55,y=1.45,w=12.2,h=4.05):
    headers=s['headers'];rows=s['rows'];n=len(headers)
    shape=slide.shapes.add_table(len(rows)+1,n,Inches(x),Inches(y),Inches(w),Inches(h));tab=shape.table
    widths=s['widths'] or [1/n]*n
    for j,v in enumerate(widths):tab.columns[j].width=Inches(w*v/sum(widths))
    fontsize=19 if n<5 else 17.5
    while True:
        heights=[max(estimate_lines(str(v),fontsize,w*widths[j]/sum(widths)-.25) for j,v in enumerate(row))*fontsize*1.1+11 for row in [headers]+rows]
        if sum(heights)<=h*72 or fontsize<=15:break
        fontsize-=.5
    extra=max(0,h*72-sum(heights))/len(heights)
    for i,hh in enumerate(heights):tab.rows[i].height=Pt(hh+extra)
    for i,row in enumerate([headers]+rows):
        for j,val in enumerate(row):
            cell=tab.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(255,255,255)
            cell.margin_left=Inches(.12);cell.margin_right=Inches(.10);cell.margin_top=Inches(.07);cell.margin_bottom=Inches(.05)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE;tf=cell.text_frame;tf.clear();tf.word_wrap=True
            p=tf.paragraphs[0];p.line_spacing=1.10
            run=p.add_run();run.text=str(val);setfont(run,fontsize,bold=(i==0))
            tcPr=cell._tc.get_or_add_tcPr()
            for edge in ['lnL','lnR','lnT','lnB']:
                ln=OxmlElement('a:'+edge);ln.set('w','6350')
                fill=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr');rgb.set('val','A0A0A0');fill.append(rgb);ln.append(fill);tcPr.append(ln)
    return shape

def arrow(slide,x1,y1,x2,y2):
    sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=RGBColor(0,0,0);sh.line.width=Pt(1.2)
    ln=sh._element.spPr.get_or_add_ln();end=OxmlElement('a:tailEnd');end.set('type','triangle');ln.append(end)

def outline(slide,label,x,y,w,h):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.background();sh.line.color.rgb=RGBColor(0,0,0);sh.line.width=Pt(1)
    tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.10);tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(label.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=PP_ALIGN.CENTER;p.line_spacing=1.1
        r=p.add_run();r.text=line;setfont(r,20)
    return sh

def schematic(slide,name):
    if name=='direction':
        outline(slide,'已有工作\n通用模型估计 V(S)',.65,1.9,3.25,1.25)
        outline(slide,'新增计算\n背景边际 → M、Γ',4.9,1.9,3.25,1.25)
        outline(slide,'论文产物\n字段风险＋支撑背景',9.15,1.9,3.25,1.25)
        arrow(slide,4.0,2.52,4.75,2.52);arrow(slide,8.25,2.52,9.,2.52)
        tx(slide,'从“给定集合能推断多少”推进到“每个字段在什么背景下带来多少新增风险”。',.75,3.75,11.9,.8,22)
    elif name=='estimator':
        outline(slide,'重建主干\n在训练数据上预学习',.65,1.55,3.0,1.15)
        outline(slide,'随机主干\n固定随机初始化',.65,3.2,3.0,1.15)
        outline(slide,'种子内拼接\n重建＋随机＋原始字段\n逐集合求 ridge 读出',4.55,2.02,3.75,1.8)
        outline(slide,'三个种子分别求解\n预测平均后得到 V̂(S)',9.15,2.20,3.48,1.45)
        arrow(slide,3.75,2.12,4.42,2.6);arrow(slide,3.75,3.77,4.42,3.23);arrow(slide,8.4,2.92,9.0,2.92)
        tx(slide,'共同设置：标准差下限、训练范围特征截断、五折选 λ、训练标签范围预测截断。',.75,4.65,11.9,.72,21)

def render(d):
    prs=Presentation();prs.slide_width=12192000;prs.slide_height=6858000
    prs.core_properties.title=d['title'];prs.core_properties.subject='组会进展汇报';prs.core_properties.author=''
    records=[]
    for idx,s in enumerate(d['slides'],1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);fill=slide.background.fill;fill.solid();fill.fore_color.rgb=RGBColor(255,255,255)
        size=32
        while estimate_lines(s['title'],size,12.25)>1 and size>26:size-=1
        tx(slide,s['title'],.55,.32,12.25,.72,size,True,BLUE)
        kind=s['kind']
        if kind=='intro':
            tx(slide,s['lead'],.65,1.35,12.0,.6,20)
            bulletbox(slide,s['bullets'],.85,2.45,11.65,3.7,24,27)
        elif kind=='text':
            start=1.40
            if s['lead']:tx(slide,s['lead'],.65,1.32,12.,.7,22,True);start=2.25
            bulletbox(slide,s['bullets'],.75,start,11.9,6.88-start,23,22)
        elif kind=='table':
            table(slide,s,y=1.45,h=4.00)
            bulletbox(slide,s['bullets'],.65,5.72,12.,1.32,19,8)
        elif kind=='diagram':
            schematic(slide,s['diagram']);bulletbox(slide,s['bullets'],.65,5.72,12.,1.25,19,8)
        elif kind=='formula':
            tx(slide,s['lead'],.70,1.47,11.9,2.20,25)
            bulletbox(slide,s['bullets'],.75,4.25,11.9,2.45,21,15)
        elif kind=='side':
            picture(slide,s['images'][0],.55,1.28,7.30,5.65)
            bulletbox(slide,s['bullets'],8.20,1.55,4.53,5.30,21,22)
        else:
            n=len(s['images']);gap=.25;w=(12.22-(n-1)*gap)/n
            for j,im in enumerate(s['images']):picture(slide,im,.55+j*(w+gap),1.17,w,4.28)
            bulletbox(slide,s['bullets'],.65,5.68,12.,1.35,19,8)
        folders=list(dict.fromkeys(re.sub('DNN_Aggresvation','',x.split('/')[0]) for x in s['sources'] if not x.startswith('reports/')))
        source='对应：'+('、'.join(folders) if folders else d['scope'])
        if s['images']:
            source+='　｜　'+('原始数据 Matplotlib 单独绘图' if any(x.get('derived') for x in s['images']) else '实验原图'+('（含子图裁剪）' if any(x['crop'] for x in s['images']) else ''))
        tx(slide,source,.55,7.22,11.5,.20,8.5,False,GRAY)
        tx(slide,f'{idx}/{len(d["slides"])}',12.12,7.20,.65,.25,10,False,GRAY,PP_ALIGN.RIGHT)
        notes=[f'第 {d["no"]} 次组会 · 第 {idx} 页',s['title'],'',*s['bullets'],'',*s['notes'],'','来源：',*s['sources']]
        if s['images']:notes+=['','图像处理：']+[x['source']+('；原合并图标签无法干净分离，读取原CSV按原图计算口径，用Matplotlib单独画出。' if x.get('derived') else ('；原始图片直接嵌入，无重绘。' if not x['crop'] else f'；原始图片直接嵌入，仅PPT显示裁剪 {x["crop"]}，可恢复。')) for x in s['images']]
        slide.notes_slide.notes_text_frame.text='\n'.join(notes)
        records.append({'slide':idx,'title':s['title'],'sources':s['sources'],'images':s['images'],'notes':s['notes']})
    folder=OUT/f'{d["no"]:02d}_{d["title"]}'/'V1';folder.mkdir(parents=True,exist_ok=True)
    path=folder/f'{d["no"]:02d}_{d["title"]}_V1.pptx';prs.save(path)
    (folder/'图表来源.md').write_text('\n'.join([
      f'# 第 {d["no"]} 次组会 V1：图表来源','',
      '白底、普通文字与原始实验图。原图裁剪在 PowerPoint 内完成，可恢复；未使用旧 HTML 的重绘图。只有第5次汇报的123号运行状态子图因面板标签相邻，使用原CSV按原计算口径用Matplotlib单独绘制，已逐页标注。','',
      *sum(([f'## 第 {r["slide"]} 页：{r["title"]}','',*[f'- `{x}`' for x in r['sources']],
             *[f'- 图片处理：{x["source"]}；'+('Matplotlib 从原CSV单独重绘，详见备注和脚本' if x.get('derived') else ('直接使用原图' if not x['crop'] else f'PPT 可恢复裁剪 {x["crop"]}')) for x in r['images']], ''] for r in records),[])
      ]),encoding='utf-8')
    return {'number':d['no'],'title':d['title'],'file':str(path.relative_to(OUT)),'slides':records,'page_count':len(records)}

def main():
    exec((HERE/'ppt_content.py').read_text(),globals())
    reports=[render(d) for d in DECKS]
    source_paths=sorted({s for d in DECKS for page in d['slides'] for s in page['sources']})
    sources=[]
    for path in source_paths:
        p=ROOT/path
        if not p.is_file():raise FileNotFoundError(path)
        sources.append({'path':path,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    (HERE/'ppt_manifest.json').write_text(json.dumps(reports,ensure_ascii=False,indent=2),encoding='utf-8')
    (HERE/'source_manifest.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')
    for r in reports:print(r['file'],r['page_count'],'页')
    print('Original image placements:',sum(len(s['images']) for d in DECKS for s in d['slides']))

if __name__=='__main__':main()
