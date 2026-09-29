"""Build the reviewable report, provenance index, diagram and editable tables."""
from pathlib import Path
import csv
import hashlib
import json
import os
import re
import subprocess
from copy import deepcopy
from io import BytesIO
from urllib.parse import quote
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1]
BASE = OUT.parent
REPO = BASE.parents[1]
WORK = OUT/'工作底稿'
ASSET = OUT/'assets'
TEMPLATE = OUT/'交付模板'

def link(label, path):
    rel=os.path.relpath(path,OUT)
    return f'[{label}]({quote(rel, safe="/._-")})'

def provenance():
    rows=json.loads((WORK/'来源清单.json').read_text())
    lines=['# 来源索引与核对口径','',
           '原始文件均保留，压缩包只在本地解压。以下链接可回到原文及按段落/表格/页码提取的工作文本。正文中的 S、R、W 编号分别指提供材料、近期研究和公开核对来源。','',
           '## 1. 用户提供材料','',
           '| 编号 | 原文 | 提取文本 |','|---|---|---|']
    for r in rows:
        p=BASE/r['source']
        lines.append(f"| {r['id']} | {link(p.name,p)} | {link('本地提取',OUT/r['extracted'])} |")
    refs=[
        ('R01','近期 IEEE 研究稿：风险定义、估计器、结果与局限',REPO/'paper/claude/V6_ieee/main.tex'),
        ('R02','V6 版本说明及验证边界',REPO/'paper/claude/V6_ieee/README.md'),
        ('R03','2026-09-25 论文主线重构与完善计划',REPO/'reports/20260925_V4论文主线重构与完善计划.md'),
        ('R04','早期冻结实验协议，仅用于辨认历史口径',REPO/'paper/FINAL_PROTOCOL.md'),
    ]
    lines+=['','## 2. 近期研究','', '| 编号 | 本地文件 | 核对说明 |','|---|---|---|']
    for key,title,p in refs:
        if not p.exists():raise FileNotFoundError(p)
        lines.append(f'| {key} | {link(title,p)} | 本次只读；不代表论文全部实验已重新运行 |')
    lines+=['',
        'R01 的具体分文件：`sec_front.tex` 为定义与方法，`sec_setup.tex` 为数据口径，`sec_application.tex` 为处置应用，`sec_discussion.tex` 为边界，`tables/tab_risk.tex` 为本文引用的组合风险结果。',
        '', '## 3. 公开的正式来源核对','',
        '| 编号 | 来源 | 本次用途 |','|---|---|---|',
        '| W01 | [国家能源局：能源行业数据分类分级指南（2026年版）](https://www.nea.gov.cn/20260630/3b456b6f83144abcb989110dcfba4d7d/c.html) | 核对正式版本、生效日期、识别条件与衍生数据条款 |',
        '| W02 | [国家能源局：指南答记者问](https://www.nea.gov.cn/20260630/2873fa450d2e4317b33d40285b5ed576/c.html) | 核对行业识别与其他保护义务之间的边界 |',
        '| W03 | [GB/T 43697-2024 官方标准信息](https://openstd.samr.gov.cn/bzgk/std/newGbInfo?hcno=F0C385EDC38CBF277AEC021F23126ADE) | 仅核对标准名称、状态和实施日期 |',
        '| W04 | [国家能源局：电力市场信息披露基本规则发布通知](https://henb.nea.gov.cn/xxgk/zcfg/202402/t20240226_257102.html) | 核对国家层面的规则来源 |',
        '', '公开来源核对日期为 2026-09-29。这是针对关键依据的核对，不构成对全部法律法规及公司制度现行效力的穷尽检索；南网材料正文状态以包内原件为准。',
        '', '## 4. 提取与统计口径','',
        '- DOCX：按正文 XML 顺序标注 P（段落）和 T（表格），表格展开至行；不依赖自动目录的旧页码。',
        '- 论文多数公式是 MathType 图形，纯文本存在缺失；另以本地转换 PDF 物理页 75—76 目视核对了第 4.2 节的核心公式。转换页码不是原 Word 排版的固定页码。',
        '- XLSX：直接读取共享字符串和工作表 XML，保留原始单元格定位；资源表统计包含序号 1—561，数据区没有需要继承的合并单元格。',
        '- PDF：按物理页提取，引用以细则条号优先。',
        '- 原始文件 SHA-256 和大小见工作底稿中的来源清单 JSON；目录统计及全部 561 项记录均保存。',
        '- 本次没有对师兄论文、近期论文的实验或南网实际业务数据重新训练、复测。材料中实验数据统一按原文自报结果处理。',
        '', '## 5. 复现本次整理','',
        '在本目录运行以下脚本可重建提取文本、目录统计和报告。脚本只对本目录生成物和相邻解压目录写入；解压遇到同名但内容不同的文件时会停止。','',
        '```bash\npython 工作底稿/extract_sources.py\npython 工作底稿/analyze_catalog.py\npython 工作底稿/build_report.py\n```',
        '', 'PDF 阅读版由生成的 Word 本地转换，保留在报告目录中。提取文本不替代原始文件。',
    ]
    (OUT/'来源索引.md').write_text('\n'.join(lines)+'\n')
    for key,title,p in refs:
        rows.append({'id':key,'source':str(p.relative_to(REPO)), 'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    (WORK/'本次研究来源快照.json').write_text(json.dumps(rows[-4:],ensure_ascii=False,indent=2))

def diagram():
    ASSET.mkdir(exist_ok=True)
    im=Image.new('RGB',(1800,1540),'#f5f7fa');d=ImageDraw.Draw(im)
    fontfile='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
    fonts={s:ImageFont.truetype(fontfile,s,index=2) for s in [28,32,36,44,52]}
    d.text((100,58),'规则与推断证据协同分类分级体系',font=fonts[52],fill='#142c48')
    d.text((100,140),'南网制度确定边界　｜　既有方法获取线索　｜　近期研究补充组合证据',font=fonts[32],fill='#486174')
    boxes=[
        ('01  规则与资产对齐','条款及版本 · 业务分类 · 数据资源、字段与加工血缘','#285680'),
        ('02  基线分类分级','影响对象与程度 · 当前批准级别 · 待核实条件与规则冲突','#285680'),
        ('03  定义开放共享场景','接收方与用途 · 披露时点及粒度 · 已知背景 · 受保护目标','#285680'),
        ('04  关联发现与组合证据','关联图 / 常见模式 → 集合能力 / 字段边际 → 独立复核','#087e82'),
        ('05  形成两类治理建议','等级复核建议　＋　实际保护方案、剩余风险与使用效果','#915824'),
        ('06  审批、执行与持续复评','结果固化目录 · 策略落实 · 证据留痕 · 变更触发复评','#915824'),
    ]
    for i,(title,body,color) in enumerate(boxes):
        y=235+i*185
        d.rounded_rectangle((100,y,1585,y+145),radius=18,fill='white',outline='#dce4eb',width=2)
        d.rounded_rectangle((100,y,116,y+145),radius=5,fill=color)
        d.text((145,y+18),title,font=fonts[44],fill=color)
        d.text((145,y+87),body,font=fonts[32],fill='#2c4254')
        if i<5:
            x=830;d.line((x,y+148,x,y+177),fill='#91a6b7',width=4)
            d.polygon([(x-10,y+166),(x+10,y+166),(x,y+179)],fill='#91a6b7')
    d.line([(1598,1275),(1690,1275),(1690,480),(1600,480)],fill='#087e82',width=5)
    d.polygon([(1610,469),(1610,491),(1594,480)],fill='#087e82')
    for k,ch in enumerate('变化触发复评'):
        d.text((1720,700+k*45),ch,font=fonts[32],fill='#087e82')
    d.text((100,1400),'贯穿全程：规则来源、场景边界、证据范围、审批责任和实际执行可追溯',font=fonts[36],fill='#142c48')
    d.text((100,1465),'本图为整合方案；推断风险分数不直接等于正式数据级别。',font=fonts[28],fill='#627a8b')
    im.save(ASSET/'体系流程.png',optimize=True)

def tables():
    TEMPLATE.mkdir(exist_ok=True)
    schemas={
      '01_分类分级主表':['对象ID','资源目录序号','系统或接口','数据对象名称','对象类型','字段清单','业务分类路径','辅助分类标签','来源与血缘','数据规模及计量口径','时间空间精度','覆盖范围','当前批准级别','一般重要核心身份','识别认定状态','影响对象','影响程度与理由','依据条款及版本','规则冲突ID','建议复核级别','责任部门','审批记录ID','生效时间','复评触发条件'],
      '02_场景与授权表':['场景ID','对象ID','接收主体ID','市场注册角色及范围','数据空间注册入驻状态','资源授权状态','用途','评估时点','实际披露时刻','时间空间粒度','既有背景B','额外背景范围H及可得依据','拟提供集合A或产品版本','受保护目标Y','目标保护依据','背景预算K与覆盖范围','历史提供记录','披露义务','合同授权与有效期','可用保护措施'],
      '03_组合风险证据表':['证据ID','场景ID','字段ID','目标ID','数据版本与哈希','样本时间范围','训练验证测试划分','背景基线能力V_B','实际提供集合能力V_BA','无额外背景边际M0','最大边际估计M_K','背景放大Gamma','已复核最大边际','关键背景及来源','扫描覆盖范围','经验模型范围','独立复核方法','不确定性及未覆盖范围','直接暴露或强代理情况','危害分析证据','对应检查项','复核人','证据日期'],
      '04_处置与变更表':['变更ID','场景ID','对象ID','证据ID','变更触发事件','是否申请等级复核','现有批准级别','建议级别及依据','批准级别','实际保护措施','新产品ID及处理版本','规则与披露义务满足情况','措施前后风险','业务效用与成本','剩余风险','责任人','审批人及审批记录','执行策略ID','目录同步记录','执行验证','回滚条件','下次复评时间'],
    }
    for name,headers in schemas.items():
        with (TEMPLATE/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.writer(f);w.writerow(headers)
    (TEMPLATE/'填写说明.md').write_text('''# 四张主表填写说明

CSV 使用 UTF-8 BOM，可由 Excel 导入。本次只提供表头，不填造任何南网等级、审批或实测风险结果。

- 对象 ID、场景 ID、证据 ID、变更 ID 用于关联四张表。
- 当前批准级别、建议级别与批准后级别分别填写；未知为“待核实”，不能用 0 替代。
- 规则依据记录实际条款及版本，草案须标明；模型分数不进入法定身份字段。
- V、M 等数值需同时填写数据版本、目标、场景、K、模型范围和样本划分；未评估留空并说明。
- 经验复核结果与正式统计置信区间分别描述；无充分依据不要填写“已证明安全”。
- 背景可得性、实际披露时刻与历史提供记录应有来源。K 只是分析预算。
- 每次脱敏/统计处理生成新的产品版本，原始对象的级别与产品级别分别管理。
- 模板是本次提出的项目设计，不是南网原有法定或审批表单。对接时可映射到现有系统字段。
''')

def word_report():
    # LibreOffice 6.4 cannot import this nested OMML expression reliably.
    # Preserve editable LaTeX in the source and render this one formula for Word/PDF.
    formula='U(A) = ∑ᵢ [−βΓᵢ exp(−λoᵢ) − η(exp(σoᵢ)−1) − ω max(ôᵢ−oᵢ, 0)²]'
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',44)
    bbox=font.getbbox(formula)
    equation=Image.new('RGB',(bbox[2]+30,bbox[3]-bbox[1]+40),'white')
    ImageDraw.Draw(equation).text((15,20-bbox[1]),formula,font=font,fill='black')
    equation.save(ASSET/'论文效用公式.png',dpi=(220,220))
    default_ref=subprocess.run(['pandoc','--print-default-data-file=reference.docx'],check=True,capture_output=True).stdout
    ref=Document(BytesIO(default_ref))
    sec=ref.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
    sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for s in ref.styles:
        if s.type in (1,2):
            s.font.name='Noto Sans CJK SC'
            s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    normal=ref.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.2
    normal.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',26),('Subtitle',12),('Heading 1',19),('Heading 2',15),('Heading 3',12)]:
        s=next(s for s in ref.styles if s.name==name);s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('183D5A')
        s.paragraph_format.keep_with_next=True
        s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(7)
    for sname in ['Body Text','First Paragraph']:
        if sname in ref.styles:ref.styles[sname].paragraph_format.first_line_indent=Pt(0)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    run=footer.add_run('南网分类分级体系梳理  |  ');run.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    refpath=WORK/'报告样式.docx';ref.save(refpath)
    parts=['00_南网数据分类分级一体化体系方案.md','01_师兄论文指定章节梳理.md','02_南网文件逐份解读.md','03_汇报主线与落地清单.md','来源索引.md']
    combined='---\ntitle: 南网数据分类分级体系梳理与整合方案\nsubtitle: 论文基础、南网规则与背景组合推断研究的一体化设计\nauthor: 项目讨论稿\ndate: 2026-09-29\nlang: zh-CN\ntoc-title: 目录\n---\n\n'
    combined+='''**阅读导引**

主方案：体系定位、规则接入、场景与证据、治理决定、试点及交付。

附录一：师兄论文 3.1、3.4.2 和第 4 章逐节梳理。

附录二：南网九份文件解读、分级矩阵、开放目录统计与版本差异。

附录三：七页汇报主线、项目研究内容与实施清单。

附录四：原文、近期研究与公开核对来源索引。

建议先读主方案；需要核对原文或准备专题汇报时再查阅对应附录。

'''
    for i,name in enumerate(parts):
        s=(OUT/name).read_text()
        if name=='01_师兄论文指定章节梳理.md':
            s=re.sub(r'\$\$\s*U\(A\)=.*?\$\$', '![](assets/论文效用公式.png){width=16cm}',s,count=1,flags=re.S)
        if i:s=re.sub(r'^# ',f'# 附录{["","一","二","三","四"][i]}：',s,count=1)
        combined+='\n\n'+s
    combined_path=WORK/'完整报告_合并稿.md';combined_path.write_text(combined)
    dst=OUT/'南网分类分级体系梳理_完整报告.docx'
    subprocess.run(['pandoc',str(combined_path),'-f','markdown+tex_math_dollars','--resource-path',str(OUT),'--reference-doc',str(refpath),'-o',str(dst)],check=True,cwd=OUT)
    doc=Document(dst)
    for original in list(doc.tables):
        old_rows=original._tbl.tr_lst
        ncol=max(len(row.tc_lst) for row in old_rows)
        table=doc.add_table(rows=len(old_rows),cols=ncol)
        for row,newrow in zip(old_rows,table.rows):
            for oldcell,newcell in zip(row.tc_lst,newrow.cells):
                for child in list(newcell._tc):
                    if child.tag!=qn('w:tcPr'):newcell._tc.remove(child)
                for child in oldcell:
                    if child.tag!=qn('w:tcPr'):newcell._tc.append(deepcopy(child))
        original._tbl.addprevious(table._tbl)
        original._tbl.getparent().remove(original._tbl)
        table.autofit=False
        table.style='Table'
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            line=OxmlElement('w:'+edge);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'4');line.set(qn('w:color'),'CED9E2');borders.append(line)
        table._tbl.tblPr.append(borders)
        weights={2:[.28,.72],3:[.22,.42,.36],4:[.16,.28,.28,.28],5:[.10,.18,.24,.24,.24],6:[.28,.144,.144,.144,.144,.144]}.get(ncol,[1/ncol]*ncol)
        widths=[Cm(17*v) for v in weights]
        for col,width in zip(table.columns,widths):col.width=width
        tw=table._tbl.tblPr.find(qn('w:tblW'))
        tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(Cm(17).twips))
        for ri,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if ri==0:
                repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
            for ci,cell in enumerate(row.cells):
                width=widths[ci]
                cell.width=width
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.line_spacing=1.1
                    p.paragraph_format.keep_with_next=False
                    for run in p.runs:run.font.size=Pt(9)
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8F0F5');cell._tc.get_or_add_tcPr().append(shade)
    for p in doc.paragraphs:
        if p.style.name=='Heading 1':p.paragraph_format.page_break_before=True
    doc.save(dst)
    print(dst)

if __name__=='__main__':
    provenance();diagram();tables();word_report()
